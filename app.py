import io
import os
import uuid
import shutil
import gc
import docx
import time
from pypdf import PdfReader
from fastapi import FastAPI, File, UploadFile, HTTPException
from pydantic import BaseModel
from vector_store import add_chunk, retrieve, get_context, store_document, delete_document
from generate import generate_answer
from memory import add_message, get_history, init_db
from typing import Optional
from fastapi.middleware.cors import CORSMiddleware
from contextlib import asynccontextmanager
from celery.result import AsyncResult
from pathlib import Path
from tasks import celery_app, process_video_task

MAX_FILE_SIZE = 5 * 1024 * 1024

STORAGE_DIR = Path("storage/videos")
STORAGE_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTENSIONS = {".mp4", ".mkv", ".mov", ".avi", ".webm"}
MAX_VIDEO_SIZE = 500 * 1024 * 1024

@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    yield
    
app = FastAPI(title="RAG Document Q&A", lifespan=lifespan)

@app.get("/")
def health_check():
    return {"status": "healthy", "service": "RAG API"}

origins = [
    "http://localhost:3000",
    "https://ecommerce-fullstack-design-olive.vercel.app", 
    "https://disease-prediction-model-pi.vercel.app"
]

app.add_middleware(
    CORSMiddleware,
    allow_origins= ["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

class QueryRequest(BaseModel):
    session_id: str
    question: str
    filename: Optional[str] = None

@app.post("/upload-video")
async def upload_video_endpoint(file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(status_code=400, detail="Filename is required")
    
    file_ext = Path(file.filename).suffix.lower()
    if file_ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail= f"Unsupported video format '{file_ext}'. \nAllowed: {', '.join(ALLOWED_EXTENSIONS)}"
        )
        
    unique_filename = f"{uuid.uuid4().hex}_{file.filename.lower()}"
    DEST_PATH = STORAGE_DIR / unique_filename

    try:
        with open(DEST_PATH, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to save video: {str(e)}")
    finally:
        await file.close()
        
    file_size = os.path.getsize(DEST_PATH)
    if file_size > MAX_VIDEO_SIZE:
        os.remove(DEST_PATH)
        raise HTTPException(status_code=413, detail=f"File size is not allowed to exceed 500MB.")
    
    task = process_video_task.delay(str(DEST_PATH), file.filename.lower())
    return {
        "status": "queued",
        "task_id": task.id,
        "filename": file.filename.lower(),
        "check_status_url": f"/tasks/{task.id}"
    }
    
@app.get("/tasks/{task_id}")
def get_task_status(task_id: str):
    try:
        uuid.UUID(task_id)
    except ValueError:
        raise HTTPException(status_code=400, detail=f"Invalid task_id format. Must be a valid UUID")
    
    task_result = AsyncResult(task_id, app=celery_app)
    
    response = {
        "task_id": task_id,
        "state": task_result.state,
    }
    
    if task_result.state == "PENDING":
        response["progress"] = 0
        response["message"] = "Task is queued and waiting for a worker."
    
    elif task_result.state == "PROCESSING":
        meta = task_result.info if isinstance(task_result.info, dict) else {}
        response["progress"] = meta.get("progress", 25)                        #Looks up progress in self.update_state, otherwise just shows 25% done
        response["message"] = meta.get("step", "Processing video.....")
        
    elif task_result.state == "SUCCESS":
            meta = task_result.info if isinstance(task_result.info, dict) else {}
            response["progress"] = 100                        
            response["result"] = task_result.result
    
    elif task_result.state == "FAILURE":
        response["progress"] = 0
        response["error"] = str(task_result.info)
    
    return response
    
    
@app.post("/upload")
async def upload_document(file: UploadFile = File(...)):
    content = await file.read()
    
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File size is too large. MAX allowed is 5MB."
        )
    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="Filename is required."
        )  
    
    file_name = file.filename.lower()

    if file_name.endswith(".docx"):
        doc = docx.Document(io.BytesIO(content))
        text = "\n".join([para.text for para in doc.paragraphs])

    elif file_name.endswith(".txt"):
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("cp1252", errors="replace")

    elif file_name.endswith(".pdf"):
        pdf = PdfReader(io.BytesIO(content))
        text = "\n".join([page.extract_text() or "" for page in pdf.pages])

    else:
        raise HTTPException(
            status_code = 400,
            detail= "Unsupported file type. Try .txt, .pdf or .docx file."
        )
    
    if not text.strip():
        raise HTTPException(
            status_code=400,
            detail="Could not extract any readable text. If this is a scanned PDF, it requires OCR."
        )
    safe_filename = file.filename.lower() if file.filename else "Unknown"
    word_count, is_small = store_document(text, safe_filename)   
    
    # Reclaim temporary buffer memory immediately
    del content
    del text
    gc.collect()

    return {
        "status": "ok", 
        "word_count": word_count, 
        "mode": "full-context" if is_small else "retrieval"
    }

class DeleteRequest(BaseModel):
    filename: str

@app.post("/delete")
def delete_document_endpoint(request: DeleteRequest):
    safe_filename = request.filename.lower() if request.filename else None
    delete_document(request.filename)
    return {"status": "ok", "deleted": request.filename}

@app.post("/query")
def queryDocument(request: QueryRequest):
    
    safe_filename = request.filename.lower() if request.filename else None

    t0 = time.time()
    relevant_chunks = get_context(request.question, safe_filename)
    print(f"Retrieval took: {time.time() - t0:.2f}s")
    
    t1 = time.time()
    history = get_history(request.session_id)
    print(f"History took: {time.time() - t1:.2f}s")
    
    print(f"\n----- Debug Query ------\n")
    print(f"filename sent to get_context: {safe_filename}")
    print(f"Context type: {type(relevant_chunks)}, Length: {len(relevant_chunks) if isinstance(relevant_chunks, list) else 'N/A'}")
    if isinstance(relevant_chunks, list) and relevant_chunks:
        print(f"First chunk (first 150 chars): {relevant_chunks[0][:150]}")

    t2 = time.time()
    answer = generate_answer(request.question, relevant_chunks, history)
    print(f"Gemini took: {time.time() - t2:.2f}s")

    t3 = time.time()

    if not answer.startswith("⚠️"):
        add_message(request.session_id, "user", request.question)
        add_message(request.session_id, "assistant", answer)
        print(f"Saving history: {time.time() - t3:.2f}s")

    gc.collect()

    return {"answer": answer, "sources": relevant_chunks}

if __name__ == "__main__":
    import uvicorn
    import os
    port = int(os.environ.get("PORT", 1024))
    uvicorn.run("app:app", host="0.0.0.0", port=port, reload=False)